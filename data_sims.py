from flask import (
    Blueprint, render_template, request, jsonify,
    send_file, session, redirect, url_for
)
import pandas as pd
from functools import wraps
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from data_store import load_data  # cache global

data_sims_bp = Blueprint("data_sims", __name__)

TEMPLATE_FILE = "Template Format Pemeriksaan UPLOAD.xlsx"

# =====================
# LOGIN REQUIRED
# =====================
def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if "user" not in session:
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper

# =====================
# FILTER DATA
# =====================
def apply_filter(search_value):
    df = load_data()
    if not search_value:
        return df

    keywords = search_value.lower().split()
    angka = [k for k in keywords if k.isdigit()]
    teks = [k for k in keywords if not k.isdigit()]

    def match(row):
        cells = [str(c).lower() for c in row]

        # angka harus exact match
        for a in angka:
            if not any(c == a for c in cells):
                return False

        # teks partial OR
        if teks:
            if not any(any(t in c for c in cells) for t in teks):
                return False

        return True

    return df[df.apply(match, axis=1)]

# =====================
# PAGE DATA SIMS
# =====================
@data_sims_bp.route("/data-sims", strict_slashes=False)
@login_required
def page():
   cols = [c for c in load_data().columns if c.lower() != "no"]
   return render_template(
    "data_sims.html",
    columns=cols
)


# =====================
# API DATATABLES
# =====================
@data_sims_bp.route("/data", strict_slashes=False)
@login_required
def api():
    df = load_data()

    draw = int(request.args.get("draw", 1))
    start = int(request.args.get("start", 0))
    length = int(request.args.get("length", 10))
    search = request.args.get("search[value]", "").strip()

    filtered = apply_filter(search)

    # HAPUS KOLOM 'no' DARI TAMPILAN DATA SIMS
    if "no" in filtered.columns:
        filtered = filtered.drop(columns=["no"])

    page = filtered.iloc[start:start + length]

    return jsonify({
        "draw": draw,
        "recordsTotal": len(df),
        "recordsFiltered": len(filtered),
        "data": page.values.tolist()
    })


# =====================
# HELPER
# =====================
def div_1000(val):
    try:
        return float(val) / 1000
    except:
        return ""

# =====================
# DOWNLOAD EXCEL
# =====================
@data_sims_bp.route("/data-sims/download", strict_slashes=False)
@login_required
def download():
    search_value = request.args.get("search", "").strip()
    filtered = apply_filter(search_value)

    wb = load_workbook(TEMPLATE_FILE)
    ws = wb.active
    START_ROW = 7

    dv_metode = DataValidation(
        type="list",
        formula1='"Inspeksi melalui Open Shelter,Pemeriksaan melalui Remote Site"',
        allow_blank=True
    )
    dv_sertifikat = DataValidation(
        type="list",
        formula1='"Ada,Tidak"',
        allow_blank=True
    )
    dv_status = DataValidation(
        type="list",
        formula1='"Sesuai ISR,Tidak Sesuai Parameter Teknis,Tidak Berizin,Tidak Aktif"',
        allow_blank=True
    )

    ws.add_data_validation(dv_metode)
    ws.add_data_validation(dv_sertifikat)
    ws.add_data_validation(dv_status)

    for i, (_, r) in enumerate(filtered.iterrows()):
        row = START_ROW + i

        ws.cell(row=row, column=1).value = i + 1
        dv_metode.add(ws.cell(row=row, column=3))

        ws.cell(row=row, column=4).value = r.get("CLNT_ID", "")
        ws.cell(row=row, column=5).value = r.get("CLNT_NAME", "")
        ws.cell(row=row, column=7).value = r.get("CURR_LIC_NUM", "")
        ws.cell(row=row, column=8).value = r.get("LINK_ID", "")
        ws.cell(row=row, column=9).value = r.get("STN_NAME", "")
        ws.cell(row=row, column=10).value = r.get("STASIUN_LAWAN", "")
        ws.cell(row=row, column=11).value = r.get("SID_LONG", "")
        ws.cell(row=row, column=12).value = r.get("SID_LAT", "")
        ws.cell(row=row, column=13).value = r.get("FREQ", "")
        ws.cell(row=row, column=14).value = r.get("FREQ_PAIR", "")
        bwidth = r.get("BWIDTH", "")
        ws.cell(row=row, column=15).value = div_1000(bwidth)
        ws.cell(row=row, column=16).value = r.get("EQ_MDL", "")
        ws.cell(row=row, column=17).value = r.get("STN_NAME", "")
        ws.cell(row=row, column=18).value = r.get("STASIUN_LAWAN", "")
        ws.cell(row=row, column=19).value = r.get("LONG", "")
        ws.cell(row=row, column=20).value = r.get("LAT", "")
        ws.cell(row=row, column=21).value = r.get("FREQ", "")
        ws.cell(row=row, column=22).value = r.get("FREQ_PAIR", "")
        ws.cell(row=row, column=23).value = div_1000(bwidth)
        ws.cell(row=row, column=24).value = r.get("EQ_MDL", "")

        dv_sertifikat.add(ws.cell(row=row, column=25))
        dv_status.add(ws.cell(row=row, column=26))
        ws.cell(row=row, column=27).value = r.get("MULAI BEROPERASI", "")
        ws.cell(row=row, column=28).value = r.get("KETERANGAN", "")
        ws.cell(row=row, column=29).value = r.get("CITY", "")

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        download_name="laporan_data_sims.xlsx",
        as_attachment=True
    )
